"""Fixture test for the watcher's change path: watch.py detects a new or corrected rate card and propose.py turns it
into a valid repo update with a CHANGELOG entry and a draft post. Runs offline on a copy of the repo.

    python -m unittest scripts/test_watch.py      (needs openpyxl)

Class guarded: a Meta change that the watcher misses, or a proposed update that fails validation or loses data.
"""
import csv
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
REPO = HERE.parent
CARD = sorted((REPO / "raw").iterdir())[-1]  # Meta's archived workbooks for the current card


def make_fixture(dest, new_title_date=None, bump=1.1):
    """Copy the archived workbooks; bump Mexico's marketing list rate in every currency, optionally re-date them."""
    dest.mkdir(parents=True)
    for src in CARD.glob("*.xlsx"):
        wb = openpyxl.load_workbook(src)
        ws = wb.active
        if new_title_date:
            ws["A1"] = ws["A1"].value.replace("October 1, 2026", new_title_date)
        if src.name.endswith("_list_rates.xlsx"):
            for row in ws.iter_rows():
                if row[0].value == "Mexico":
                    row[2].value = round(row[2].value * bump, 4)
        wb.save(dest / src.name)


def copy_repo(dest):
    for name in ("data", "scripts", "raw"):
        shutil.copytree(REPO / name, dest / name, ignore=shutil.ignore_patterns("__pycache__"))
    for name in ("rates.json", "datapackage.json", "CHANGELOG.md", "CITATION.cff"):
        shutil.copy2(REPO / name, dest / name)


def run(*args):
    r = subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
    return dict(line.split("=", 1) for line in r.stdout.splitlines() if "=" in line and " " not in line.split("=")[0])


def usd(repo, market, cat):
    with open(repo / "data" / "latest" / "USD.csv", newline="", encoding="utf8") as f:
        return next(r for r in csv.DictReader(f) if r["market"] == market)[cat]


class WatcherChangePath(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.repo = self.tmp / "repo"
        copy_repo(self.repo)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def watch(self, fixture):
        work = self.tmp / "work"
        out = run(self.repo / "scripts" / "watch.py", work, "--offline", fixture, "--repo", self.repo)
        return work, out["status"]

    def test_unchanged_card_is_ok(self):
        _, status = self.watch(CARD)
        self.assertEqual(status, "ok")

    def test_new_card_becomes_a_valid_update(self):
        fixture = self.tmp / "fixture"
        make_fixture(fixture, "January 1, 2027")
        old = usd(self.repo, "Mexico", "marketing")
        work, status = self.watch(fixture)
        self.assertEqual(status, "changed")
        self.assertIn("New rate card: effective 2027-01-01", (work / "report.md").read_text(encoding="utf8"))

        out = run(self.repo / "scripts" / "propose.py", work, "--repo", self.repo, "--checked", "2026-12-01")
        self.assertEqual(out["status"], "ready")
        self.assertEqual(out["branch"], "watcher/rates-2027-01-01")
        self.assertTrue((self.repo / "data" / "rates" / "2027-01-01.csv").exists())
        self.assertTrue((self.repo / "data" / "rates" / f"{CARD.name}.csv").exists(), "old snapshot kept")
        self.assertTrue((self.repo / "raw" / "2027-01-01" / "USD_list_rates.xlsx").exists())
        self.assertNotEqual(usd(self.repo, "Mexico", "marketing"), old)
        self.assertIn('"latest_effective_date": "2027-01-01"', (self.repo / "rates.json").read_text(encoding="utf8"))
        log = (self.repo / "CHANGELOG.md").read_text(encoding="utf8")
        self.assertLess(log.index("## 2027-01-01 (current)"), log.index(f"## {CARD.name}\n"))
        self.assertIn("USD list rates changed: 1 ", log)
        body = (work / "pr-body.md").read_text(encoding="utf8")
        post = body.split("```\n")[1].split("\n```")[0]
        self.assertIn("effective 2027-01-01: 1 USD list rate changed", post)
        self.assertRegex(post, r"Mexico marketing \+(9\.9|10\.[01])%")  # +10%, after 4-decimal rounding
        self.assertLessEqual(len(post), 300)
        self.assertIn('version: "2027-01-01"', (self.repo / "CITATION.cff").read_text(encoding="utf8"))

    def test_mid_quarter_correction_keeps_the_date(self):
        fixture = self.tmp / "fixture"
        make_fixture(fixture, bump=0.9)
        work, status = self.watch(fixture)
        self.assertEqual(status, "changed")
        self.assertIn("List rates: 16 changed", (work / "report.md").read_text(encoding="utf8"))
        out = run(self.repo / "scripts" / "propose.py", work, "--repo", self.repo, "--checked", "2026-11-02")
        self.assertEqual(out["status"], "ready")
        self.assertEqual(len(list((self.repo / "data" / "rates").glob("*.csv"))), 1)
        self.assertIn(f"## {CARD.name} (correction found 2026-11-02)",
                      (self.repo / "CHANGELOG.md").read_text(encoding="utf8"))


if __name__ == "__main__":
    unittest.main()
