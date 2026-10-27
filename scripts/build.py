"""Build the tidy dataset from Meta's raw rate-card workbooks.

    python scripts/build.py --raw DIR [--country-map FILE.md] [--checked YYYY-MM-DD] [--out .] [--archive-raw]

DIR holds <CUR>_list_rates.xlsx and <CUR>_volume_tiers.xlsx (see fetch.py). Values are Meta's own, per
currency, never converted. Cleaning only: header newlines, HTML entities, "n/a" to absent, scientific
notation to plain decimal strings, merged-cell market names filled down.
"""
import argparse
import csv
import datetime as dt
import hashlib
import html
import json
import re
import shutil
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import meta_source as ms  # noqa: E402

CATS = ["marketing", "utility", "authentication", "authentication_international", "service"]
LIST_COLS = {"marketing": 2, "utility": 3, "authentication": 4, "authentication_international": 5, "service": 6}
TIER_BLOCKS = {"utility": 2, "authentication": 7, "authentication_international": 12}
SOURCE = ms.PAGE


def dec(v):
    """Cell value to a plain decimal string, or None for empty / n/a / '--'."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if s.lower() in ("", "n/a", "--", "-"):
            return None
        d = Decimal(s.replace(",", ""))
    else:
        d = Decimal(repr(v))
    return format(d.normalize(), "f") if d != 0 else "0"


def clean_market(v):
    return html.unescape(str(v)).replace("\n", " ").strip()


def effective_date(title):
    m = re.search(r"effective ([A-Z][a-z]+ \d{1,2}, \d{4})", title or "")
    if not m:
        raise ValueError(f"no effective date in title: {title!r}")
    return dt.datetime.strptime(m.group(1), "%B %d, %Y").date().isoformat()


def read_list(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    eff = effective_date(rows[0][0])
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Market")
    out = []
    for r in rows[hdr + 1:]:
        if not r[0]:
            continue
        out.append((clean_market(r[0]), r[1], {c: dec(r[i]) for c, i in LIST_COLS.items()}))
    return eff, out


def read_tiers(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    eff = effective_date(rows[0][0])
    first = next(i for i, r in enumerate(rows) if r[2] == "From") + 1
    out, market = [], None
    for r in rows[first:]:
        if r[0]:
            market = clean_market(r[0])
        if market is None or r[1] is None:
            continue
        for cat, c0 in TIER_BLOCKS.items():
            frm, to, rtype, rate, pct = (r[c0 + k] for k in range(5))
            if dec(rate) is None:
                continue
            out.append((market, cat, dec(frm), dec(to), str(rtype).strip().lower().replace(" ", "_"),
                        dec(rate), dec(pct)))
    return eff, out


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def read_csv(path):
    with open(path, newline="", encoding="utf8") as f:
        return list(csv.DictReader(f))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fields(spec):
    return {"fields": [{"name": n, "type": t, "description": d} for n, t, d in spec]}


def write_datapackage(out, index, curs):
    """datapackage.json (Frictionless) describing every file; regenerated on each build."""
    rate_f = [("effective_date", "date", "First day the rates apply (Meta Messaging account timezone)"),
              ("currency", "string", "ISO 4217 code; Meta's own published value, never converted"),
              ("market", "string", "Market name as written on Meta's rate card"),
              ("category", "string", "marketing, utility, authentication, authentication_international or service"),
              ("rate", "string", "Decimal string, per delivered message, excluding VAT; absent when Meta lists n/a")]
    tier_f = rate_f[:4] + [("tier_from", "string", "First message of the tier in the calendar month"),
                           ("tier_to", "string", "Last message of the tier; empty = open ended (Meta writes --)"),
                           ("rate_type", "string", "list_rate or tier_rate"),
                           ("rate", "string", "Decimal string per message"),
                           ("pct_vs_list", "string", "Decimal fraction vs the list rate, e.g. -0.05")]
    wide_f = [("currency", "string", ""), ("effective_date", "date", ""), ("checked_on", "date", "Day this repo last compared with Meta"),
              ("source", "string", "Meta page the values come from"), ("market", "string", ""),
              ("type", "string", "country, region or other"), ("iso_codes", "string", "Space separated ISO 3166-1 alpha-2 members"),
              ("calling_code", "string", "Space separated country calling codes"),
              *[(c, "string", "Rate, empty when not published") for c in CATS]]
    cm_f = [("iso2", "string", "ISO 3166-1 alpha-2"), ("country", "string", "Name as Meta lists it"),
            ("market", "string", "Rate-card market this country is priced under"),
            ("market_type", "string", "country or region"), ("calling_code", "string", ""),
            ("extra_area_codes", "string", "Shared-code area codes Meta lists, e.g. 809 829 849")]
    res = []
    for v in index["versions"]:
        d = v["effective_date"]
        res.append({"name": f"rates-{d}", "path": v["rates"], "format": "csv", "mediatype": "text/csv",
                    "schema": fields(rate_f), "hash": "sha256:" + v["rates_sha256"]})
        res.append({"name": f"tiers-{d}", "path": v["tiers"], "format": "csv", "mediatype": "text/csv",
                    "schema": fields(tier_f), "hash": "sha256:" + v["tiers_sha256"]})
    for c in curs:
        res.append({"name": f"latest-{c.lower()}", "path": f"data/latest/{c}.csv", "format": "csv",
                    "mediatype": "text/csv", "schema": fields(wide_f)})
    res.append({"name": "country-market", "path": "data/country_market.csv", "format": "csv",
                "mediatype": "text/csv", "schema": fields(cm_f)})
    pkg = {"name": "whatsapp-api-rates", "title": "WhatsApp Business Platform per-message rates",
           "description": "Meta's published WhatsApp Business Platform rate cards in 16 currencies, as tidy "
                          "CSV with volume tiers and a country to market table. Not affiliated with Meta.",
           "version": index["latest_effective_date"], "licenses": [{"name": "CC-BY-4.0",
                                                                      "path": "https://creativecommons.org/licenses/by/4.0/"}],
           "contributors": [{"title": "Tobachat", "role": "author"}],
           "sources": [{"title": "Meta, Pricing on the WhatsApp Business Platform", "path": SOURCE}],
           "resources": res}
    (out / "datapackage.json").write_text(json.dumps(pkg, indent=2) + chr(10), encoding="utf8")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True)
    ap.add_argument("--country-map")
    ap.add_argument("--checked", default=dt.date.today().isoformat())
    ap.add_argument("--out", default=str(Path(__file__).parent.parent))
    ap.add_argument("--archive-raw", action="store_true")
    a = ap.parse_args()
    raw, out = Path(a.raw), Path(a.out)
    data = out / "data"

    lists, tiers = {}, {}
    for p in sorted(raw.glob("*_list_rates.xlsx")):
        cur = p.name[:3]
        eff, rows = read_list(p)
        eff2, trows = read_tiers(raw / f"{cur}_volume_tiers.xlsx")
        if eff != eff2:
            raise SystemExit(f"{cur}: list ({eff}) and tier ({eff2}) dates differ")
        if any(c != cur for _, c, _ in rows):
            raise SystemExit(f"{cur}: unexpected currency column")
        lists[cur], tiers[cur] = (eff, rows), (eff, trows)
    effs = {e for e, _ in lists.values()}
    if len(effs) != 1:
        raise SystemExit(f"currencies disagree on effective date: {effs}")
    eff = effs.pop()
    curs = sorted(lists)

    cm_path = data / "country_market.csv"
    if a.country_map:
        crow, declared = ms.parse_country_map(Path(a.country_map).read_text(encoding="utf8"))
        crow.sort(key=lambda r: r[0])
        write_csv(cm_path, ["iso2", "country", "market", "market_type", "calling_code", "extra_area_codes"], crow)
        bad = {m: (n, sum(1 for r in crow if r[2] == m)) for m, n in declared.items()
               if m != "Other" and n != sum(1 for r in crow if r[2] == m)}
        if bad:
            print("NOTE: declared vs listed member counts differ:", bad)
    cmap = read_csv(cm_path)
    by_market = {}
    for r in cmap:
        by_market.setdefault(r["market"], []).append(r)

    long_rows = [(eff, cur, m, cat, v) for cur in curs for m, _, vals in lists[cur][1]
                 for cat, v in vals.items() if v is not None]
    write_csv(data / "rates" / f"{eff}.csv", ["effective_date", "currency", "market", "category", "rate"], long_rows)
    trow = [(eff, cur, *r) for cur in curs for r in tiers[cur][1]]
    write_csv(data / "tiers" / f"{eff}.csv",
              ["effective_date", "currency", "market", "category", "tier_from", "tier_to", "rate_type", "rate",
               "pct_vs_list"], trow)

    newest = max(p.stem for p in (data / "rates").glob("*.csv"))
    if newest == eff:
        for cur in curs:
            rows = []
            for m, _, vals in lists[cur][1]:
                members = by_market.get(m, [])
                if m == "Other":
                    mtype = "other"
                elif len(members) == 1 and members[0]["country"] == m:
                    mtype = "country"
                else:
                    mtype = "region"
                codes = sorted({r["calling_code"] for r in members}, key=int)
                rows.append((cur, eff, a.checked, SOURCE, m, mtype, " ".join(r["iso2"] for r in members),
                             " ".join(codes), *[vals[c] or "" for c in CATS]))
            write_csv(data / "latest" / f"{cur}.csv",
                      ["currency", "effective_date", "checked_on", "source", "market", "type", "iso_codes",
                       "calling_code", *CATS], rows)

    versions = []
    for p in sorted((data / "rates").glob("*.csv")):
        t = data / "tiers" / p.name
        versions.append({"effective_date": p.stem, "rates": f"data/rates/{p.name}", "rates_sha256": sha(p),
                         "tiers": f"data/tiers/{p.name}", "tiers_sha256": sha(t)})
    index = {"name": "whatsapp-api-rates", "publisher": "Tobachat", "license": "CC-BY-4.0", "source": SOURCE,
             "latest_effective_date": newest, "checked_on": a.checked,
             "currencies": [{"code": c, "latest": f"data/latest/{c}.csv"} for c in curs],
             "country_market": "data/country_market.csv", "versions": versions}
    (out / "rates.json").write_text(json.dumps(index, indent=2) + "\n", encoding="utf8")
    write_datapackage(out, index, curs)
    if a.archive_raw:
        dest = out / "raw" / eff
        dest.mkdir(parents=True, exist_ok=True)
        for p in raw.glob("*.xlsx"):
            shutil.copy2(p, dest / p.name)
    print(f"built {eff}: {len(curs)} currencies, {len(long_rows)} rate rows, {len(trow)} tier rows")


if __name__ == "__main__":
    main()
